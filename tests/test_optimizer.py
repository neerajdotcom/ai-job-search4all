"""
tests/test_optimizer.py — Unit tests for resume optimization and DOCX paragraph patching.
"""

import pytest
from pathlib import Path
from docx import Document
from optimizer.resume_optimizer import (
    _find_paragraph_index,
    _replace_paragraph_text,
    extract_protected_metrics,
    patch_docx,
)


def test_find_paragraph_index():
    doc = Document()
    doc.add_paragraph("John Doe")
    doc.add_paragraph("Professional Summary")
    doc.add_paragraph("Experience")
    doc.add_paragraph("Acme Corp")

    assert _find_paragraph_index(doc, "Professional Summary") == 1
    assert _find_paragraph_index(doc, "Experience") == 2
    assert _find_paragraph_index(doc, "Nonexistent Heading") is None


def test_replace_paragraph_text():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Original ").bold = True
    p.add_run("Text")
    assert p.text == "Original Text"

    _replace_paragraph_text(p, "New Updated Content")
    assert p.text == "New Updated Content"


def test_extract_protected_metrics():
    text = "Achieved 40% growth and $10M ARR with 3x efficiency."
    metrics = extract_protected_metrics(text)
    assert "40%" in metrics
    assert "$10M" in metrics
    assert "3x" in metrics


def test_patch_docx_dry(tmp_path, sample_profile):
    """Create a temporary DOCX and verify paragraph patching without throwing."""
    doc = Document()
    doc.add_paragraph("John Doe")
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph("Original summary line.")
    doc.add_heading("Core Competencies", level=1)
    doc.add_paragraph("Python, Go")
    doc.add_heading("Professional Experience", level=1)
    
    p_acme = doc.add_paragraph()
    p_acme.add_run("Acme Corp — Senior Software Engineer").bold = True
    doc.add_paragraph("Led backend development achieving 20% latency drop.")

    doc_path = tmp_path / "test_resume.docx"
    doc.save(str(doc_path))

    tailoring_result = {
        "optimized_summary": "Tailored executive summary.",
        "optimized_competencies": ["Python", "Go", "PostgreSQL", "Cloud"],
        "optimized_bullets": {
            "acme_corp": ["Tailored bullet for Acme Corp with 20% latency drop."],
        },
    }

    ats_path, review_path, warnings = patch_docx(
        str(doc_path),
        tailoring_result,
        sample_profile,
        company_name="TestCo",
        title="Senior Engineer",
    )
    assert Path(review_path).exists()
    assert Path(ats_path).exists()


def test_parse_resume_sections_unwrapping():
    """Verify that hard-wrapped lines from PDF extracts are merged into full sentences, not 1-word bullets."""
    from optimizer.resume_generator import parse_resume_sections, assemble_tailored_resume_data, generate_tailored_pdf

    raw_wrapped_resume = """
    Neeraj Banerjee
    neerajbanerjee@gmail.com | Kolkata, India | linkedin.com/in/neerajbanerjee
    
    PROFESSIONAL SUMMARY
    Results-driven Project Manager with 5 years experience across gaming and iGaming pipelines.
    
    PROFESSIONAL EXPERIENCE
    Project/Delivery Manager | Arrise Solutions (October 2023–Present)
    • delivery
    rate
    and
    ensuring
    alignment
    with
    strategic
    objectives.
    • Risk Management: Proactively identified and mitigated 20+ high-impact risks,
    contributing to a 30% decrease in project delays.
    
    ACADEMICS
    Master in Business Administration | Pune University (2021-2023)
    
    CERTIFICATIONS
    Certified Scrum Master (CSM)
    """

    parsed = parse_resume_sections(raw_wrapped_resume)
    assert parsed["name"] == "Neeraj Banerjee"
    assert len(parsed["experience"]) >= 1
    
    bullets = parsed["experience"][0]["bullets"]
    # Check that unwrapping joined the fragments into full sentences
    assert any("delivery rate and ensuring alignment" in b for b in bullets)
    assert not any(len(b.split()) == 1 for b in bullets)

    # Test PDF generation does not throw and produces a valid PDF stream
    resume_data = assemble_tailored_resume_data(
        resume_text=raw_wrapped_resume,
        profile_dict={"name": "Neeraj Banerjee", "title": "Project Manager", "location": "Kolkata, India"},
        tailored_pack=None,
        job={"title": "Project Manager", "company": "Ubisoft"},
    )
    pdf_buffer = generate_tailored_pdf(resume_data)
    assert pdf_buffer.getvalue().startswith(b"%PDF")
    assert len(pdf_buffer.getvalue()) > 1000

